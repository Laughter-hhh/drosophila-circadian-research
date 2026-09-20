import csv
import sys
import tempfile
import unittest
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import openpyxl

from scripts.extract_published_sc_clock_channel_rhythms import (
    _validate_output_paths,
    extract,
    sha256_file,
)


HEADERS = (
    "Gene", "cluster", "F24", "phase.F24", "F24.p.value", "JTK_pvalue",
    "JTK_BH.Q", "MAX", "MIN", "Max.Min", "cycling.", "Condition",
)


class PublishedSingleCellChannelRhythmTests(unittest.TestCase):
    def _fixture(self, root: Path, rows_ld=None, rows_dd=None, candidates=None):
        candidate_path = root / "candidates.csv"
        candidate_rows = candidates or [
            {"gene_symbol": "alpha", "priority_class": "leak"},
            {"gene_symbol": "beta", "priority_class": "gated"},
            {"gene_symbol": "missing", "priority_class": "leak"},
        ]
        with candidate_path.open("w", newline="", encoding="utf-8") as handle:
            writer = csv.DictWriter(handle, fieldnames=("gene_symbol", "priority_class"))
            writer.writeheader()
            writer.writerows(candidate_rows)

        workbook = openpyxl.Workbook()
        ld = workbook.active
        ld.title = "LD condition"
        ld.append(HEADERS)
        ld_rows = rows_ld or [["alpha", "2:s_LNv", 0.8, 4, 0.01, 0.02, 0.03, 3, 0.5, 6, "HC_cycler", "LD"]]
        for row in ld_rows:
            ld.append(row)
        dd = workbook.create_sheet("DD condition")
        dd.append(HEADERS)
        dd_rows = rows_dd or [["beta", "3:DN1a", 0.6, 8, 0.03, 0.04, 0.08, 2, 1, 2, "LC_cycler", "DD"]]
        for row in dd_rows:
            dd.append(row)
        info = workbook.create_sheet("Info")
        info.append(["Rhythm rule: illustrative test criterion"])
        xlsx_path = root / "supplement.xlsx"
        workbook.save(xlsx_path)
        workbook.close()
        return candidate_path, xlsx_path

    def test_real_gse157504_published_table_extraction(self):
        candidate_path = ROOT / "validation" / "public-data" / "GSE22308_candidate_genes.csv"
        supplement_path = ROOT / "validation" / "public-data" / "GSE157504_published_rhythmic_genes_supp1.xlsx"
        candidates, details, report = extract(
            candidate_path,
            supplement_path,
            "https://elifesciences.org/articles/63056",
            ROOT,
            expected_candidate_sha256="ee46819993f11916145bdab8f13e35f73895986b4cdb0d3fd9b89c6dd1f0e7ce",
            expected_supplement_sha256="d3553deb0290bb56d73ab7713ec6b73f5749dc9d466dcdb5063b91cdef5c804b",
        )
        self.assertEqual(len(candidates), 15)
        self.assertEqual(len(details), 21)
        self.assertEqual(report["n_supplement_rows_scanned"], {"LD": 2243, "DD": 1359})
        self.assertEqual(report["n_candidates_with_author_HC_rhythm_rows"], 9)
        detail_counts = Counter(row["candidate"] for row in details)
        self.assertEqual(detail_counts["Irk1"], 6)
        self.assertEqual(detail_counts["Shab"], 3)
        self.assertEqual(
            report["candidate_rows_without_HC_call"],
            ["Shal", "para", "sei", "Irk2", "Ork1", "trpl"],
        )
        self.assertIn("does not establish channel protein", report["inference_warning"])

    def test_synthetic_extract_preserves_author_class_and_no_call_limit(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            candidate_path, xlsx_path = self._fixture(root)
            summary, details, report = extract(
                candidate_path, xlsx_path, "https://example.org/supplement", root
            )
            self.assertEqual(len(summary), 3)
            self.assertEqual(len(details), 2)
            self.assertEqual(report["n_candidates_with_author_HC_rhythm_rows"], 1)
            by_gene = {row["candidate"]: row for row in summary}
            self.assertEqual(by_gene["alpha"]["published_rhythm_status"], "author_reported_HC_cycler")
            self.assertEqual(by_gene["beta"]["published_rhythm_status"], "not_listed_in_author_HC_cycler_table")
            self.assertEqual(details[1]["condition"], "DD")
            self.assertIn("not evidence", report["inference_warning"])
            self.assertEqual(sha256_file(candidate_path), report["inputs"]["candidate_list"]["sha256"])

    def test_hash_mismatch_is_rejected(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            candidate_path, xlsx_path = self._fixture(root)
            with self.assertRaisesRegex(ValueError, "SHA-256 mismatch"):
                extract(
                    candidate_path, xlsx_path, "https://example.org/supplement", root,
                    expected_candidate_sha256="0" * 64,
                )

    def test_duplicate_candidates_are_rejected(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            duplicates = [
                {"gene_symbol": "alpha", "priority_class": "leak"},
                {"gene_symbol": "alpha", "priority_class": "gated"},
            ]
            candidate_path, xlsx_path = self._fixture(root, candidates=duplicates)
            with self.assertRaisesRegex(ValueError, "duplicate gene_symbol"):
                extract(candidate_path, xlsx_path, "https://example.org/supplement", root)

    def test_nonfinite_supplement_value_is_rejected(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            bad_row = ["alpha", "2:s_LNv", "NaN", 4, 0.01, 0.02, 0.03, 3, 0.5, 6, "HC_cycler", "LD"]
            candidate_path, xlsx_path = self._fixture(root, rows_ld=[bad_row])
            with self.assertRaisesRegex(ValueError, "non-finite F24"):
                extract(candidate_path, xlsx_path, "https://example.org/supplement", root)

    def test_output_paths_cannot_overwrite_inputs_or_escape_root(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            input_path = root / "input.csv"
            input_path.write_text("x\n", encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "overwrite an input"):
                _validate_output_paths({"report": input_path}, (input_path,), root)
            with self.assertRaisesRegex(ValueError, "inside provenance root"):
                _validate_output_paths({"report": root.parent / "outside.json"}, (input_path,), root)


if __name__ == "__main__":
    unittest.main()
